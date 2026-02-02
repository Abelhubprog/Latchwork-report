from docx import Document
from docx.shared import Pt
import re

def detect_heading(text):
    """Detect heading level by content patterns"""
    text = text.strip()
    
    # Title patterns
    if text in ['Latchwork', 'Technical Documentation']:
        return 'title'
    
    # Numbered sections like "1. Introduction", "2.1 Subsection"
    if re.match(r'^\d+\.\s+\w', text):
        return 'h1'
    if re.match(r'^\d+\.\d+\s+\w', text):
        return 'h2'
    if re.match(r'^\d+\.\d+\.\d+\s+\w', text):
        return 'h3'
    
    # Table of Contents header
    if text.lower() in ['table of contents', 'contents']:
        return 'h1'
    
    # Reference section patterns
    if text in ['API Reference', 'Error Codes', 'Glossary', 'Architecture Overview']:
        return 'h2'
    
    # Tutorial patterns
    if re.match(r'^(Getting Started|Creating|Building|Deploying|Configuring)', text):
        return 'h2'
    
    # Explanation patterns
    if text in ['Understanding the Slot Model', 'Security Architecture', 'Payment Flow']:
        return 'h2'
    
    # Common section headers
    if text in ['Introduction', 'Tutorials', 'How-to Guides', 'Reference', 'Explanation']:
        return 'h1'
    
    return 'p'

def clean_text(text):
    """Clean text for HTML"""
    return text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

def make_id(text):
    """Create a valid HTML id from text"""
    return re.sub(r'[^\w\s-]', '', text).strip().replace(' ', '-').lower()[:50]

doc = Document('Latchwork_Technical_Documentation.docx')

# First pass: identify all headings and create TOC
headings = []
toc_entries = []
section_map = {}

for para in doc.paragraphs:
    text = para.text.strip()
    if not text:
        continue
    
    level = detect_heading(text)
    if level != 'p':
        sid = make_id(text)
        base_sid = sid
        counter = 1
        while sid in section_map:
            sid = f"{base_sid}-{counter}"
            counter += 1
        section_map[sid] = text
        
        if level == 'title':
            headings.append((level, text, None))
        else:
            headings.append((level, text, sid))
            toc_entries.append((level, text, sid))

# Table mapping based on document structure
table_sections = [
    ('1. introduction to latchwork', 0),  # Agent Types table
    ('4.1 agent configuration schema', 1),  # Agent Config Schema
    ('3.2 how to set up micropayments with x402', 2),  # x402 Config
    ('4.2 api endpoints', 3),  # Main API Endpoints
    ('4.2 api endpoints', 4),  # Payment API (after main endpoints)
    ('4.3 state machine transitions', 5),  # State Transitions
    ('4.4 error codes and troubleshooting', 6),  # Error Codes
    ('4.4 error codes and troubleshooting', 7),  # Troubleshooting (after error codes)
]

html_parts = []

html_parts.append('''<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Latchwork Technical Documentation</title>
    <style>
        :root {
            --primary: #2563eb;
            --primary-dark: #1d4ed8;
            --secondary: #64748b;
            --bg: #ffffff;
            --bg-alt: #f8fafc;
            --text: #1e293b;
            --text-light: #64748b;
            --border: #e2e8f0;
            --code-bg: #f1f5f9;
        }
        
        * { box-sizing: border-box; margin: 0; padding: 0; }
        
        body {
            font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, 'Helvetica Neue', Arial, sans-serif;
            line-height: 1.7;
            color: var(--text);
            background: var(--bg);
        }
        
        .container {
            max-width: 1400px;
            margin: 0 auto;
            display: flex;
        }
        
        .sidebar {
            width: 300px;
            background: var(--bg-alt);
            border-right: 1px solid var(--border);
            padding: 2rem 1.5rem;
            position: fixed;
            height: 100vh;
            overflow-y: auto;
            top: 0;
        }
        
        .sidebar-title {
            font-size: 1.1rem;
            font-weight: 600;
            color: var(--primary);
            margin-bottom: 0.5rem;
        }
        
        .sidebar-subtitle {
            font-size: 0.85rem;
            color: var(--text-light);
            margin-bottom: 1.5rem;
            padding-bottom: 1rem;
            border-bottom: 1px solid var(--border);
        }
        
        .toc-title {
            font-size: 0.75rem;
            text-transform: uppercase;
            letter-spacing: 0.1em;
            color: var(--text-light);
            margin-bottom: 1rem;
            font-weight: 600;
        }
        
        .toc { list-style: none; }
        .toc li { margin: 0.25rem 0; }
        
        .toc a {
            display: block;
            color: var(--text);
            text-decoration: none;
            padding: 0.4rem 0.5rem;
            border-radius: 0.375rem;
            font-size: 0.9rem;
            transition: all 0.15s;
        }
        
        .toc a:hover {
            background: rgba(37, 99, 235, 0.1);
            color: var(--primary);
        }
        
        .toc-h1 { font-weight: 600; }
        .toc-h2 { padding-left: 1rem !important; font-size: 0.85rem !important; }
        .toc-h3 { padding-left: 2rem !important; font-size: 0.8rem !important; color: var(--text-light); }
        
        .main {
            margin-left: 300px;
            flex: 1;
            padding: 3rem 4rem;
            max-width: 900px;
        }
        
        @media (max-width: 1100px) {
            .sidebar { display: none; }
            .main { margin-left: 0; padding: 2rem; }
        }
        
        .report-header {
            text-align: center;
            padding: 2rem 0 3rem;
            border-bottom: 2px solid var(--border);
            margin-bottom: 3rem;
        }
        
        .report-header h1 {
            font-size: 2.75rem;
            color: var(--primary);
            margin-bottom: 0.5rem;
            font-weight: 700;
            letter-spacing: -0.02em;
        }
        
        .report-header .subtitle {
            font-size: 1.35rem;
            color: var(--secondary);
            font-style: italic;
            margin-bottom: 1.5rem;
        }
        
        .report-header .meta {
            color: var(--text-light);
            font-size: 0.95rem;
            line-height: 1.6;
        }
        
        .report-header .author {
            font-weight: 600;
            color: var(--text);
        }
        
        .tags {
            margin-top: 1.5rem;
            display: flex;
            gap: 0.5rem;
            justify-content: center;
            flex-wrap: wrap;
        }
        
        .tag {
            background: var(--primary);
            color: white;
            padding: 0.35rem 1rem;
            border-radius: 9999px;
            font-size: 0.7rem;
            text-transform: uppercase;
            letter-spacing: 0.08em;
            font-weight: 600;
        }
        
        .main h1 {
            font-size: 1.875rem;
            margin: 3rem 0 1.25rem;
            color: var(--primary-dark);
            font-weight: 700;
            letter-spacing: -0.01em;
        }
        
        .main h2 {
            font-size: 1.4rem;
            margin: 2.5rem 0 1rem;
            color: var(--primary);
            font-weight: 600;
            padding-bottom: 0.5rem;
            border-bottom: 1px solid var(--border);
        }
        
        .main h3 {
            font-size: 1.15rem;
            margin: 2rem 0 0.75rem;
            color: var(--text);
            font-weight: 600;
        }
        
        .main h4 {
            font-size: 1.05rem;
            margin: 1.5rem 0 0.5rem;
            color: var(--text);
            font-weight: 600;
        }
        
        .main p {
            margin-bottom: 1.1rem;
            text-align: left;
        }
        
        .main ul, .main ol {
            margin: 1rem 0 1rem 1.5rem;
        }
        
        .main li { margin: 0.5rem 0; }
        
        .main pre {
            background: var(--code-bg);
            padding: 1.25rem;
            border-radius: 0.5rem;
            overflow-x: auto;
            border: 1px solid var(--border);
            margin: 1.25rem 0;
        }
        
        .main code {
            font-family: 'SF Mono', Monaco, 'Cascadia Code', Consolas, monospace;
            font-size: 0.9em;
        }
        
        .main pre code {
            background: none;
            padding: 0;
            font-size: 0.85rem;
            line-height: 1.6;
        }
        
        .main p code {
            background: var(--code-bg);
            padding: 0.15rem 0.4rem;
            border-radius: 0.25rem;
            color: var(--primary-dark);
        }
        
        .main table {
            width: 100%;
            border-collapse: collapse;
            margin: 1.5rem 0;
            font-size: 0.95rem;
            box-shadow: 0 1px 3px rgba(0,0,0,0.05);
        }
        
        .main th, .main td {
            padding: 0.875rem 1rem;
            text-align: left;
            border: 1px solid var(--border);
        }
        
        .main th {
            background: var(--bg-alt);
            font-weight: 600;
            color: var(--primary-dark);
            font-size: 0.85rem;
            text-transform: uppercase;
            letter-spacing: 0.03em;
        }
        
        .main tr:nth-child(even) { background: #fafbfc; }
        
        .info-box {
            background: #eff6ff;
            border-left: 4px solid var(--primary);
            padding: 1.25rem 1.5rem;
            margin: 1.5rem 0;
            border-radius: 0 0.5rem 0.5rem 0;
        }
        
        .info-box-title {
            font-weight: 600;
            color: var(--primary-dark);
            margin-bottom: 0.5rem;
            font-size: 0.9rem;
            text-transform: uppercase;
            letter-spacing: 0.03em;
        }
        
        .warning-box {
            background: #fffbeb;
            border-left-color: #f59e0b;
        }
        .warning-box .info-box-title { color: #b45309; }
        
        .tip-box {
            background: #f0fdf4;
            border-left-color: #22c55e;
        }
        .tip-box .info-box-title { color: #15803d; }
        
        .report-footer {
            margin-top: 5rem;
            padding: 2rem;
            border-top: 1px solid var(--border);
            text-align: center;
            color: var(--text-light);
            font-size: 0.875rem;
            background: var(--bg-alt);
            border-radius: 0.5rem;
        }
        
        @media print {
            .sidebar { display: none; }
            .main { margin-left: 0; max-width: none; }
        }
        
        a { color: var(--primary); }
        a:hover { color: var(--primary-dark); }
        
        h1[id], h2[id], h3[id] { scroll-margin-top: 2rem; }
    </style>
</head>
<body>
    <div class="container">
        <nav class="sidebar">
            <div class="sidebar-title">Latchwork</div>
            <div class="sidebar-subtitle">Technical Documentation</div>
            <div class="toc-title">Contents</div>
            <ul class="toc">
''')

# Generate TOC
for level, text, sid in toc_entries:
    toc_class = f"toc-{level}"
    html_parts.append(f'                <li class="{toc_class}"><a href="#{sid}">{clean_text(text)}</a></li>\n')

html_parts.append('''            </ul>
        </nav>
        
        <main class="main">
''')

# Header
html_parts.append('''            <header class="report-header">
                <h1>Latchwork</h1>
                <div class="subtitle">Technical Documentation</div>
                <div class="meta">
                    <span class="author">Abel Ngeno</span> — Technical Author | Full-stack Developer | Health Informatics Specialist<br>
                    August 2025
                </div>
                <div class="tags">
                    <span class="tag">Tutorials</span>
                    <span class="tag">How-to Guides</span>
                    <span class="tag">Reference</span>
                    <span class="tag">Explanation</span>
                </div>
            </header>
''')

def process_table(table):
    """Convert a docx table to HTML"""
    html = '<table>\n'
    for i, row in enumerate(table.rows):
        html += '  <tr>\n'
        for cell in row.cells:
            tag = 'th' if i == 0 else 'td'
            cell_text = clean_text(cell.text.strip())
            html += f'    <{tag}>{cell_text}</{tag}>\n'
        html += '  </tr>\n'
    html += '</table>\n'
    return html

# Track which tables have been used
used_tables = set()
tables_by_section = {}
for section, tidx in table_sections:
    if section not in tables_by_section:
        tables_by_section[section] = []
    tables_by_section[section].append(tidx)

# Generate main content
content_parts = []
in_list = False

for para in doc.paragraphs:
    text = para.text.strip()
    if not text:
        continue
    
    level = detect_heading(text)
    text_lower = text.lower()
    
    # Skip title/header elements
    if level == 'title' or text in ['Latchwork', 'Technical Documentation']:
        continue
    if text.startswith('Tutorials |') or text.startswith('Author:'):
        continue
    if text in ['August 2025']:
        continue
    
    # Handle headings
    if level == 'h1':
        if in_list:
            content_parts.append('</ul>')
            in_list = False
        sid = make_id(text)
        content_parts.append(f'<h1 id="{sid}">{clean_text(text)}</h1>')
        # Insert tables for this section
        if text_lower in tables_by_section:
            for tidx in tables_by_section[text_lower]:
                if tidx not in used_tables:
                    content_parts.append(process_table(doc.tables[tidx]))
                    used_tables.add(tidx)
        continue
    
    if level == 'h2':
        if in_list:
            content_parts.append('</ul>')
            in_list = False
        sid = make_id(text)
        content_parts.append(f'<h2 id="{sid}">{clean_text(text)}</h2>')
        # Insert tables for this section
        if text_lower in tables_by_section:
            for tidx in tables_by_section[text_lower]:
                if tidx not in used_tables:
                    content_parts.append(process_table(doc.tables[tidx]))
                    used_tables.add(tidx)
        continue
    
    if level == 'h3':
        if in_list:
            content_parts.append('</ul>')
            in_list = False
        sid = make_id(text)
        content_parts.append(f'<h3 id="{sid}">{clean_text(text)}</h3>')
        continue
    
    # Handle bullet lists
    if text.startswith('- ') or text.startswith('• '):
        if not in_list:
            content_parts.append('<ul>')
            in_list = True
        item_text = clean_text(text[2:])
        content_parts.append(f'<li>{item_text}</li>')
        continue
    elif in_list and not (text.startswith('- ') or text.startswith('• ')):
        content_parts.append('</ul>')
        in_list = False
    
    # Regular paragraphs
    lowered = text.lower()
    if lowered.startswith('note:'):
        content_parts.append(f'<div class="info-box"><div class="info-box-title">Note</div><p>{clean_text(text[5:].strip())}</p></div>')
    elif lowered.startswith('tip:'):
        content_parts.append(f'<div class="info-box tip-box"><div class="info-box-title">Tip</div><p>{clean_text(text[4:].strip())}</p></div>')
    elif lowered.startswith('warning:'):
        content_parts.append(f'<div class="info-box warning-box"><div class="info-box-title">Warning</div><p>{clean_text(text[8:].strip())}</p></div>')
    else:
        content_parts.append(f'<p>{clean_text(text)}</p>')

if in_list:
    content_parts.append('</ul>')

# Add any remaining tables at the end
for i in range(len(doc.tables)):
    if i not in used_tables:
        content_parts.append(process_table(doc.tables[i]))
        used_tables.add(i)

html_parts.extend(content_parts)

# Footer
html_parts.append('''
            <footer class="report-footer">
                <p><strong>Latchwork Technical Documentation</strong></p>
                <p>&copy; 2025 Abel Ngeno — Technical Author | Full-stack Developer | Health Informatics Specialist</p>
            </footer>
        </main>
    </div>
</body>
</html>
''')

# Write the file
full_html = ''.join(html_parts)
with open('index.html', 'w', encoding='utf-8') as f:
    f.write(full_html)

print(f"HTML report generated: index.html ({len(full_html)} characters)")
print(f"TOC entries: {len(toc_entries)}")
print(f"Tables included: {len(used_tables)}/{len(doc.tables)}")
print(f"\nDocument structure detected (first 20):")
for level, text, sid in headings[:20]:
    print(f"  [{level}] {text[:60]}")
